"""Excel 导入导出（REQ-050）与评论/@提醒（REQ-100）的验收测试。"""
from __future__ import annotations

import base64
import io

import pytest
from openpyxl import Workbook, load_workbook

from app import xlsx
from app.db import get_conn

from .conftest import cell_at, grant


# --------------------------------------------------------------------------- 工具


def make_xlsx(rows: list, sheet_name: str = "数据") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    ws["A1"].font = ws["A1"].font.copy(bold=True, color="FFD32F2F")
    ws["A1"].fill = ws["A1"].fill.copy(fill_type="solid", fgColor="FFEAF1FD")
    ws["A1"].alignment = ws["A1"].alignment.copy(horizontal="center")
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")


# ------------------------------------------------------------------- REQ-050 导出


def test_export_xlsx_returns_valid_file(client, admin_headers, make_sheet):
    sheet_id = make_sheet(rows=4, cols=3)
    client.put(
        f"/api/sheets/{sheet_id}/cells",
        json={
            "updates": [
                {"row": 0, "col": 0, "value": "名称"},
                {"row": 1, "col": 0, "value": "钢笔"},
                {"row": 1, "col": 1, "value": "12.5"},
                {"row": 2, "col": 1, "value": "=SUM(B2:B2)"},
            ]
        },
        headers=admin_headers,
    )
    client.put(
        f"/api/sheets/{sheet_id}/formats",
        json={"updates": [{"row": 0, "col": 0, "format": {"bold": True, "bg": "#eaf1fd", "align": "center"}}]},
        headers=admin_headers,
    )

    response = client.post(f"/api/sheets/{sheet_id}/xlsx/export", json={}, headers=admin_headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    # 中文文件名走 RFC 5987 百分号编码，响应头保持 ASCII
    assert "%" in response.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws["A1"].value == "名称"
    assert ws["A2"].value == "钢笔"
    assert ws["B2"].value == 12.5  # 纯数字按数值写入
    assert ws["B3"].value == "=SUM(B2:B2)"  # 公式原样保留，交由 Excel 计算
    assert ws["A1"].font.bold is True
    assert ws["A1"].alignment.horizontal == "center"
    assert ws["A1"].fill.fgColor.rgb.endswith("EAF1FD")
    wb.close()


def test_export_requires_authentication(client, make_sheet):
    sheet_id = make_sheet()
    assert client.post(f"/api/sheets/{sheet_id}/xlsx/export", json={}).status_code == 401


def test_export_allowed_for_viewer(client, admin_headers, make_sheet, as_user):
    sheet_id = make_sheet(rows=2, cols=2)
    headers = as_user("excel-viewer")  # 先登录（开发模式自动注册）再授权
    grant(client, admin_headers, sheet_id, "excel-viewer", "viewer")
    assert client.post(f"/api/sheets/{sheet_id}/xlsx/export", json={}, headers=headers).status_code == 200


def test_export_empty_sheet_is_rejected(client, admin_headers, make_sheet):
    """没有任何行列的表格导出应给出可读错误而不是 500。"""
    sheet_id = make_sheet(rows=1, cols=1)
    conn = get_conn()
    # 结构表才是行列的真实来源：清掉它即得到"空表"
    conn.execute("DELETE FROM structure WHERE sheet_id = ?", (sheet_id,))
    conn.commit()
    from app import crdt

    crdt.invalidate_structure_cache(sheet_id)
    response = client.post(f"/api/sheets/{sheet_id}/xlsx/export", json={}, headers=admin_headers)
    assert response.status_code == 400
    assert "空" in response.json()["error"]


# ------------------------------------------------------------------- REQ-050 导入


def test_parse_xlsx_lists_sheets(client, admin_headers, make_sheet):
    wb = Workbook()
    first = wb.active
    first.title = "一月"
    first.append(["a", "b"])
    wb.create_sheet("二月")
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()

    sheet_id = make_sheet()
    response = client.post(
        f"/api/sheets/{sheet_id}/xlsx/parse",
        json={"data": b64(buffer.getvalue())},
        headers=admin_headers,
    )
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["sheet_names"] == ["一月", "二月"]
    assert data["sheet_name"] == "一月"
    assert data["rows"] == 1 and data["cols"] == 2


def test_import_xlsx_replaces_content_and_formats(client, admin_headers, make_sheet):
    sheet_id = make_sheet(rows=3, cols=3)
    payload = make_xlsx([["项目", "金额"], ["差旅", 320], ["办公", 128.5]])
    response = client.post(
        f"/api/sheets/{sheet_id}/xlsx/import",
        json={"data": b64(payload)},
        headers=admin_headers,
    )
    assert response.status_code == 200, response.text
    result = response.json()["data"]
    assert result["rows"] == 3 and result["cols"] == 2
    assert cell_at(result, 0, 0) == "项目"
    assert cell_at(result, 1, 1) == "320"
    assert cell_at(result, 2, 1) == "128.5"
    assert result["formats"]["0:0"]["bold"] is True
    assert result["formats"]["0:0"]["align"] == "center"
    # 导入是整体替换语义
    again = client.post(
        f"/api/sheets/{sheet_id}/xlsx/import",
        json={"data": b64(make_xlsx([["只有一行"]]))},
        headers=admin_headers,
    )
    assert again.status_code == 200
    assert set(again.json()["data"]["cells"].values()) == {"只有一行"}


def test_import_specific_sheet_by_name(client, admin_headers, make_sheet):
    wb = Workbook()
    first = wb.active
    first.title = "首页"
    first.append(["first"])
    second = wb.create_sheet("明细")
    second.append(["detail-row"])
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()

    sheet_id = make_sheet()
    response = client.post(
        f"/api/sheets/{sheet_id}/xlsx/import",
        json={"data": b64(buffer.getvalue()), "sheet_name": "明细"},
        headers=admin_headers,
    )
    assert response.status_code == 200, response.text
    assert response.json()["data"]["cells"]["0:0"] == "detail-row"


def test_import_rejects_bad_content(client, admin_headers, make_sheet):
    sheet_id = make_sheet()
    not_excel = b64(b"this is definitely not an xlsx")
    response = client.post(
        f"/api/sheets/{sheet_id}/xlsx/import", json={"data": not_excel}, headers=admin_headers
    )
    assert response.status_code == 400
    assert "xlsx" in response.json()["error"].lower() or "无法读取" in response.json()["error"]

    bad_base64 = client.post(
        f"/api/sheets/{sheet_id}/xlsx/import", json={"data": "!!!!"}, headers=admin_headers
    )
    assert bad_base64.status_code == 400


def test_import_requires_editor_role(client, admin_headers, make_sheet, as_user):
    sheet_id = make_sheet()
    headers = as_user("xlsx-viewer")
    grant(client, admin_headers, sheet_id, "xlsx-viewer", "viewer")
    response = client.post(
        f"/api/sheets/{sheet_id}/xlsx/import",
        json={"data": b64(make_xlsx([["x"]]))},
        headers=headers,
    )
    assert response.status_code == 403


def test_import_creates_version_and_can_restore(client, admin_headers, make_sheet):
    sheet_id = make_sheet(rows=2, cols=2)
    client.put(
        f"/api/sheets/{sheet_id}/cells",
        json={"updates": [{"row": 0, "col": 0, "value": "原始"}]},
        headers=admin_headers,
    )
    before = client.get(f"/api/sheets/{sheet_id}/versions", headers=admin_headers).json()["data"][0]["id"]
    client.post(
        f"/api/sheets/{sheet_id}/xlsx/import",
        json={"data": b64(make_xlsx([["导入后"]]))},
        headers=admin_headers,
    )
    current = client.get(f"/api/sheets/{sheet_id}", headers=admin_headers).json()["data"]
    assert current["cells"]["0:0"] == "导入后"
    restored = client.post(f"/api/sheets/{sheet_id}/versions/{before}/restore", headers=admin_headers)
    assert restored.status_code == 200, restored.text
    assert restored.json()["data"]["cells"]["0:0"] == "原始"


def test_xlsx_module_roundtrip_unit():
    """不经 HTTP 的纯函数往返：值/公式/样式。"""
    data = xlsx.export_workbook(
        "往返",
        3,
        2,
        {"0:0": "文字", "1:0": "=SUM(A1:A1)", "2:1": "42"},
        {"0:0": {"bold": True, "bg": "#ffeeee", "align": "right"}},
        frozen_rows=1,
        frozen_cols=1,
    )
    parsed = xlsx.import_workbook(data)
    assert parsed["sheet_name"] == "往返"
    assert parsed["cells"]["2:1"] == "42"
    assert parsed["cells"]["1:0"] == "=SUM(A1:A1)"
    assert parsed["formats"]["0:0"]["bold"] is True
    assert parsed["formats"]["0:0"]["align"] == "right"


def test_xlsx_rejects_oversize_file(monkeypatch):
    from app import config

    monkeypatch.setattr(config, "MAX_XLSX_BYTES", 10)
    with pytest.raises(xlsx.ExcelError):
        xlsx.import_workbook(make_xlsx([["too big"]]))


# ------------------------------------------------------------------- REQ-100 评论


def test_comment_crud_and_mentions(client, admin_headers, make_sheet, as_user):
    sheet_id = make_sheet(rows=3, cols=3)
    ann = as_user("ann")
    as_user("bob")
    grant(client, admin_headers, sheet_id, "ann", "editor")
    grant(client, admin_headers, sheet_id, "bob", "viewer")
    base = client.get(f"/api/sheets/{sheet_id}", headers=admin_headers).json()["data"]
    row_id, col_id = base["rowIds"][1], base["colIds"][2]

    created = client.post(
        f"/api/sheets/{sheet_id}/comments",
        json={"row_id": row_id, "col_id": col_id, "body": "@bob 这里的金额请复核，@ann 已同步"},
        headers=ann,
    )
    assert created.status_code == 200, created.text
    data = created.json()["data"]
    assert data["comment"]["author"] == "ann"
    assert data["comment"]["mentions"] == ["bob", "ann"]
    assert data["counts"][f"{row_id}:{col_id}"] == 1

    listed = client.get(f"/api/sheets/{sheet_id}/comments", headers=ann).json()["data"]
    assert len(listed["comments"]) == 1
    assert {m["username"] for m in listed["members"]} == {"admin", "ann", "bob"}

    resolved = client.post(
        f"/api/sheets/{sheet_id}/comments/{data['comment']['id']}/resolve",
        json={"resolved": True},
        headers=ann,
    )
    assert resolved.json()["data"]["comment"]["resolved"] is True
    assert resolved.json()["data"]["counts"] == {}

    deleted = client.delete(f"/api/sheets/{sheet_id}/comments/{data['comment']['id']}", headers=ann)
    assert deleted.status_code == 200
    assert client.get(f"/api/sheets/{sheet_id}/comments", headers=ann).json()["data"]["comments"] == []


def test_mention_only_known_members_are_recorded(client, admin_headers, make_sheet, as_user):
    sheet_id = make_sheet()
    carol = as_user("carol")
    grant(client, admin_headers, sheet_id, "carol", "editor")
    base = client.get(f"/api/sheets/{sheet_id}", headers=admin_headers).json()["data"]
    created = client.post(
        f"/api/sheets/{sheet_id}/comments",
        json={
            "row_id": base["rowIds"][0],
            "col_id": base["colIds"][0],
            "body": "联系 @nobody@example.com 和 @admin 处理",
        },
        headers=carol,
    )
    assert created.status_code == 200, created.text
    mentions = created.json()["data"]["mentions"]
    assert "admin" in mentions
    assert all("nobody" not in m for m in mentions)


def test_viewer_cannot_comment_but_can_read(client, admin_headers, make_sheet, as_user):
    sheet_id = make_sheet()
    reader = as_user("reader")
    grant(client, admin_headers, sheet_id, "reader", "viewer")
    base = client.get(f"/api/sheets/{sheet_id}", headers=admin_headers).json()["data"]
    created = client.post(
        f"/api/sheets/{sheet_id}/comments",
        json={"row_id": base["rowIds"][0], "col_id": base["colIds"][0], "body": "我能不能评论"},
        headers=reader,
    )
    assert created.status_code == 403
    assert client.get(f"/api/sheets/{sheet_id}/comments", headers=reader).status_code == 200


def test_only_author_or_editor_can_delete_comment(client, admin_headers, make_sheet, as_user):
    sheet_id = make_sheet()
    dan, eve = as_user("dan"), as_user("eve")
    grant(client, admin_headers, sheet_id, "dan", "editor")
    grant(client, admin_headers, sheet_id, "eve", "editor")
    base = client.get(f"/api/sheets/{sheet_id}", headers=admin_headers).json()["data"]
    comment = client.post(
        f"/api/sheets/{sheet_id}/comments",
        json={"row_id": base["rowIds"][0], "col_id": base["colIds"][0], "body": "dan 的评论"},
        headers=dan,
    ).json()["data"]["comment"]
    # eve 是 editor → 可删除他人评论
    assert client.delete(f"/api/sheets/{sheet_id}/comments/{comment['id']}", headers=eve).status_code == 200


def test_comment_on_missing_cell_is_rejected(client, admin_headers, make_sheet):
    sheet_id = make_sheet()
    response = client.post(
        f"/api/sheets/{sheet_id}/comments",
        json={"row_id": "不存在", "col_id": "也不存在", "body": "x"},
        headers=admin_headers,
    )
    assert response.status_code == 400
    assert "不存在" in response.json()["error"]


def test_comment_requires_authentication(client, make_sheet):
    sheet_id = make_sheet()
    assert client.get(f"/api/sheets/{sheet_id}/comments").status_code == 401
    assert client.post(f"/api/sheets/{sheet_id}/comments", json={"body": "x"}).status_code == 401


def test_comments_follow_element_id_after_insert(client, admin_headers, make_sheet):
    """评论绑定稳定元素 ID：插入行后仍指向同一行，不随索引漂移。"""
    sheet_id = make_sheet(rows=4, cols=2)
    base = client.get(f"/api/sheets/{sheet_id}", headers=admin_headers).json()["data"]
    target_row = base["rowIds"][2]
    client.post(
        f"/api/sheets/{sheet_id}/comments",
        json={"row_id": target_row, "col_id": base["colIds"][0], "body": "挂在第 3 行"},
        headers=admin_headers,
    )
    inserted = client.post(
        f"/api/sheets/{sheet_id}/ops", json={"kind": "insert_row", "payload": {"at": 0}}, headers=admin_headers
    )
    assert inserted.status_code == 200, inserted.text
    comments = client.get(f"/api/sheets/{sheet_id}/comments", headers=admin_headers).json()["data"]["comments"]
    assert len(comments) == 1
    assert comments[0]["row_id"] == target_row


def test_parse_mentions_unit():
    from app.store import parse_mentions

    assert parse_mentions("@ann @bob 看一下") == ["ann", "bob"]
    assert parse_mentions("没有提及") == []
    assert parse_mentions("@ann 和 @ann") == ["ann"]
    assert parse_mentions("@ghost 不存在", ["ann", "bob"]) == []
    assert parse_mentions("@ann 存在", ["ann"]) == ["ann"]


# ------------------------------------------------------------------ 删除表格


def test_delete_sheet_removes_everything(client, admin_headers, make_sheet, as_user):
    """删除表格要连带清理单元格、版本、评论与权限（不留孤儿数据）。"""
    from app.db import get_conn

    sheet_id = make_sheet(rows=3, cols=3)
    editor = as_user("cleaner")
    grant(client, admin_headers, sheet_id, "cleaner", "editor")
    client.put(
        f"/api/sheets/{sheet_id}/cells",
        json={"updates": [{"row": 0, "col": 0, "value": "会被删掉"}]},
        headers=admin_headers,
    )
    base = client.get(f"/api/sheets/{sheet_id}", headers=admin_headers).json()["data"]
    client.post(
        f"/api/sheets/{sheet_id}/comments",
        json={"row_id": base["rowIds"][0], "col_id": base["colIds"][0], "body": "删表前的评论"},
        headers=admin_headers,
    )

    assert client.delete(f"/api/sheets/{sheet_id}", headers=editor).status_code == 403  # 非所有者
    assert client.delete(f"/api/sheets/{sheet_id}", headers=admin_headers).status_code == 200
    assert client.get(f"/api/sheets/{sheet_id}", headers=admin_headers).status_code == 404
    assert sheet_id not in [s["id"] for s in client.get("/api/sheets", headers=admin_headers).json()["data"]]

    conn = get_conn()
    for table in ("permissions", "cells", "structure", "versions", "ops", "comments"):
        left = conn.execute(f"SELECT COUNT(*) AS n FROM {table} WHERE sheet_id = ?", (sheet_id,)).fetchone()["n"]
        assert left == 0, f"{table} 还残留 {left} 行"
